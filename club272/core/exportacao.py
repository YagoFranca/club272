"""
Geração de planilhas para relatórios.

Escreve Excel (.xlsx) formatado ou CSV, escolhendo pela extensão do arquivo.

Duas armadilhas motivaram este módulo:

- **O separador.** O Excel usa o separador de lista da região do Windows. Em
  pt-BR isso é `;`, não a vírgula que o Python escreve por padrão — e um CSV
  com vírgula abre com tudo amontoado numa única coluna.
- **A codificação.** Sem o BOM do `utf-8-sig`, o Excel lê os acentos errado.
"""

import csv
import locale
import os


def separador_do_sistema():
    """Separador de lista da região configurada no sistema.

    No Windows vem do registro (`Control Panel\International\sList`), que é
    exatamente de onde o Excel tira o dele — em pt-BR vale `;`. Seguir essa
    configuração é o que faz o arquivo abrir com as colunas separadas ao dar
    duplo clique; com vírgula, o Excel brasileiro amontoa tudo numa coluna só.

    Fora do Windows, cai no locale. `localeconv()` só reflete a região depois
    de um `setlocale`, que o Python não faz sozinho — por isso a chamada
    temporária, com o valor anterior restaurado logo em seguida.
    """
    try:
        import winreg

        with winreg.OpenKey(
            winreg.HKEY_CURRENT_USER, r"Control Panel\International"
        ) as chave:
            separador = winreg.QueryValueEx(chave, "sList")[0]
            if separador:
                return separador
    except (ImportError, OSError):
        pass

    anterior = locale.setlocale(locale.LC_NUMERIC)
    try:
        locale.setlocale(locale.LC_NUMERIC, "")
        virgula_decimal = locale.localeconv().get("decimal_point") == ","
    except locale.Error:
        virgula_decimal = False
    finally:
        try:
            locale.setlocale(locale.LC_NUMERIC, anterior)
        except locale.Error:
            pass
    return ";" if virgula_decimal else ","


# Identidade visual do relatório, alinhada à paleta da interface.
_COR_TITULO = "14100E"
_COR_CABECALHO = "E29464"
_COR_FAIXA = "FBF1EA"


def exportar(caminho, titulo, colunas, linhas, metadados=None, resumo=None):
    """Grava o relatório. `.xlsx` sai formatado; qualquer outra extensão, CSV.

    `metadados` e `resumo` são listas de pares (rótulo, valor) mostradas antes
    e depois da tabela.
    """
    if os.path.splitext(caminho)[1].lower() == ".xlsx":
        return _escrever_xlsx(caminho, titulo, colunas, linhas, metadados, resumo)
    return _escrever_csv(caminho, titulo, colunas, linhas, metadados, resumo)


def _escrever_csv(caminho, titulo, colunas, linhas, metadados, resumo):
    delimitador = separador_do_sistema()

    # utf-8-sig: sem o BOM o Excel abre os acentos errados.
    with open(caminho, "w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=delimitador)

        escritor.writerow([titulo])
        for rotulo, valor in (metadados or []):
            escritor.writerow([rotulo, valor])
        if metadados:
            escritor.writerow([])

        escritor.writerow(colunas)
        escritor.writerows(linhas)

        if resumo:
            escritor.writerow([])
            for rotulo, valor in resumo:
                escritor.writerow([rotulo, valor])

    return caminho


def _escrever_xlsx(caminho, titulo, colunas, linhas, metadados, resumo):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    aba = livro.active
    aba.title = "Relatório"

    ultima_coluna = max(len(colunas), 2)
    linha = 1

    # Título
    aba.cell(row=linha, column=1, value=titulo)
    aba.cell(row=linha, column=1).font = Font(size=14, bold=True, color=_COR_TITULO)
    aba.merge_cells(start_row=linha, start_column=1,
                    end_row=linha, end_column=ultima_coluna)
    aba.row_dimensions[linha].height = 24
    linha += 2

    # Metadados
    for rotulo, valor in (metadados or []):
        aba.cell(row=linha, column=1, value=rotulo).font = Font(bold=True, size=9)
        aba.cell(row=linha, column=2, value=valor).font = Font(size=9)
        linha += 1
    if metadados:
        linha += 1

    # Cabeçalho da tabela
    linha_cabecalho = linha
    borda = Side(style="thin", color="D8C9BD")
    for coluna, nome in enumerate(colunas, start=1):
        celula = aba.cell(row=linha, column=coluna, value=nome)
        celula.font = Font(bold=True, color=_COR_TITULO, size=10)
        celula.fill = PatternFill("solid", fgColor=_COR_CABECALHO)
        celula.alignment = Alignment(vertical="center")
        celula.border = Border(bottom=borda)
    aba.row_dimensions[linha].height = 20
    linha += 1

    # Dados, com faixa alternada para a leitura não se perder na horizontal
    for indice, registro in enumerate(linhas):
        for coluna, valor in enumerate(registro, start=1):
            celula = aba.cell(row=linha, column=coluna, value=valor)
            celula.font = Font(size=10)
            if indice % 2:
                celula.fill = PatternFill("solid", fgColor=_COR_FAIXA)
        linha += 1

    # Resumo, como rodapé: rótulo ocupando as colunas da esquerda e o valor
    # na última. Deixar o rótulo só na coluna A o faria ser cortado — o Excel
    # só derrama texto para o lado quando a célula vizinha está vazia, e ali
    # ela tem o valor.
    if resumo:
        linha += 1
        for rotulo, valor in resumo:
            celula = aba.cell(row=linha, column=1, value=rotulo)
            celula.font = Font(bold=True, size=10)
            celula.alignment = Alignment(horizontal="right")
            if ultima_coluna > 1:
                aba.merge_cells(start_row=linha, start_column=1,
                                end_row=linha, end_column=ultima_coluna - 1)

            celula_valor = aba.cell(row=linha, column=ultima_coluna, value=valor)
            celula_valor.font = Font(bold=True, size=10, color=_COR_TITULO)
            celula_valor.fill = PatternFill("solid", fgColor=_COR_FAIXA)
            linha += 1

    _ajustar_larguras(aba, colunas, linhas, get_column_letter)

    # Congela o cabeçalho: listas longas continuam legíveis ao rolar.
    aba.freeze_panes = aba.cell(row=linha_cabecalho + 1, column=1)
    # Filtro automático nas colunas da tabela.
    if linhas:
        aba.auto_filter.ref = (
            f"A{linha_cabecalho}:"
            f"{get_column_letter(len(colunas))}{linha_cabecalho + len(linhas)}"
        )

    livro.save(caminho)
    return caminho


def _ajustar_larguras(aba, colunas, linhas, get_column_letter):
    """Largura de cada coluna pelo conteúdo mais longo, com limites."""
    for indice, nome in enumerate(colunas):
        maior = len(str(nome))
        for registro in linhas:
            if indice < len(registro):
                maior = max(maior, len(str(registro[indice])))
        aba.column_dimensions[get_column_letter(indice + 1)].width = min(
            max(maior + 4, 12), 45
        )
